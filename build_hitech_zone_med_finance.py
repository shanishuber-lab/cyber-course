from pathlib import Path
from shutil import copy2
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\shani\Documents\course creator")
OUT = ROOT / "output" / "hitech-zone"
IMG_DIR = OUT / "images-new-courses"
IMG_DIR.mkdir(parents=True, exist_ok=True)
REFERENCE = OUT / "חבילת תוכן להייטק זון - Z-SCHOOL.docx"
FINAL = OUT / "קורסים נוספים להייטק זון - רפואה ושוק ההון - Z-SCHOOL.docx"
LOGO = ROOT / "assets" / "logos" / "לוגו צבעוני על רקע שקוף.png"
Image.MAX_IMAGE_PIXELS = None

SOURCES = {
    "medical": Path(r"C:\Users\shani\.codex\generated_images\01a018d2-1ef6-7511-9026-5aa759e0d054\exec-137d3c8f-41d3-4406-9040-ca6840461e30.png"),
    "finance": Path(r"C:\Users\shani\.codex\generated_images\01a018d2-1ef6-7511-9026-5aa759e0d054\exec-69c2b560-6b51-46ff-a369-c20ac67ac63f.png"),
}

def branded_image(src, dst):
    base = Image.open(src).convert("RGB")
    logo_rgb = Image.open(LOGO).convert("RGB")
    logo_rgb.thumbnail((2200, 1200), Image.Resampling.LANCZOS)
    mask = Image.new("L", logo_rgb.size)
    mask.putdata([255 if min(px) < 238 else 0 for px in logo_rgb.getdata()])
    bbox = mask.getbbox()
    logo_rgb, mask = logo_rgb.crop(bbox), mask.crop(bbox)
    logo = logo_rgb.convert("RGBA")
    logo.putalpha(mask)
    target_w = int(base.width * 0.145)
    ratio = target_w / logo.width
    logo = logo.resize((target_w, int(logo.height * ratio)), Image.Resampling.LANCZOS)
    badge = Image.new("RGBA", (logo.width + 36, logo.height + 24), (255, 255, 255, 235))
    badge.alpha_composite(logo, (18, 12))
    canvas = base.convert("RGBA")
    canvas.alpha_composite(badge, (28, 24))
    canvas.convert("RGB").save(dst, quality=95)

for key, src in SOURCES.items():
    branded_image(src, IMG_DIR / f"{key}-hitech-zone.jpg")

copy2(REFERENCE, FINAL)
doc = Document(FINAL)
body = doc._element.body
sectPr = body.sectPr
for child in list(body):
    if child is not sectPr:
        body.remove(child)

PURPLE, NAVY, YELLOW, CYAN, LIGHT, GRAY = "5324F5", "07152B", "FFD95A", "1AB8E8", "F4F1FF", "595F6B"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd")) or OxmlElement("w:shd")
    if shd.getparent() is None: tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def margins(cell, top=120, start=150, bottom=120, end=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar") or OxmlElement("w:tcMar")
    if tcMar.getparent() is None: tcPr.append(tcMar)
    for name, value in (("top",top),("start",start),("bottom",bottom),("end",end)):
        el = tcMar.find(qn(f"w:{name}")) or OxmlElement(f"w:{name}")
        if el.getparent() is None: tcMar.append(el)
        el.set(qn("w:w"), str(value)); el.set(qn("w:type"), "dxa")

def rtl(p):
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None: pPr.append(OxmlElement("w:bidi"))

def runfmt(r, size=11, bold=False, color=NAVY):
    r.font.name = "Arial"; r.font.size = Pt(size); r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color)
    rf = r._element.get_or_add_rPr().rFonts
    for key in ("ascii","hAnsi","cs"): rf.set(qn(f"w:{key}"), "Arial")
    r._element.get_or_add_rPr().append(OxmlElement("w:rtl"))

def para(text="", size=11, bold=False, color=NAVY, after=5, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph(); p.alignment = align
    if align == WD_ALIGN_PARAGRAPH.RIGHT: rtl(p)
    p.paragraph_format.space_after = Pt(after); p.paragraph_format.line_spacing = 1.15
    runfmt(p.add_run(text), size, bold, color); return p

def heading(text):
    p = para(text, 13, True, NAVY, 5); p.paragraph_format.space_before = Pt(6); return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet"); rtl(p); p.paragraph_format.space_after = Pt(3)
    runfmt(p.add_run(text), 10.5, False, NAVY)

def info(rows):
    t = doc.add_table(rows=len(rows), cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    for i,(label,value) in enumerate(rows):
        a,b = t.rows[i].cells; a.width=Inches(4.7); b.width=Inches(1.5)
        for c in (a,b): c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; margins(c)
        shade(b, LIGHT); rtl(a.paragraphs[0]); rtl(b.paragraphs[0])
        runfmt(a.paragraphs[0].add_run(value),10.5); runfmt(b.paragraphs[0].add_run(label),10.5,True,PURPLE)
    doc.add_paragraph()

def prices():
    t=doc.add_table(rows=1,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for c,(label,price,fill,color) in zip(t.rows[0].cells,[("מחיר רגיל","800 ₪",LIGHT,PURPLE),("מחיר הייטק זון","600 ₪",YELLOW,NAVY)]):
        c.width=Inches(3.1); shade(c,fill); margins(c,180,150,180,150)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        runfmt(p.add_run(label+"\n"),10,True,color); runfmt(p.add_run(price),18,True,color)

def product(title,subtitle,image,intro,facts,benefits,notes,source):
    doc.add_page_break(); para("Z-SCHOOL  |  הצעה להייטק זון",9,True,CYAN,2)
    para(title,24,True,PURPLE,2); para(subtitle,13,True,GRAY,10)
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(IMG_DIR/image),width=Inches(6.45)); p.paragraph_format.space_after=Pt(8)
    para(intro,11,False,NAVY,7); info(facts); heading("מה מקבלים?")
    for x in benefits: bullet(x)
    heading("כדאי לדעת")
    for x in notes: bullet(x)
    heading("אופן המימוש")
    steps=["לאחר הרכישה בהייטק זון מתקבל קוד הטבה.","יוצרים קשר עם Z-SCHOOL ומציינים את הקוד.","מקבלים קישור ייעודי, מזינים את הקוד ומשלימים את הרכישה במחיר ההטבה.","פרטי הכניסה ל-Z-Campus ולמפגש נשלחים לאחר השלמת ההרשמה."]
    for i,x in enumerate(steps,1): para(f"{i}. {x}",10.5,False,NAVY,3)
    prices(); para(f"למידע נוסף: {source}",8.5,False,GRAY,0)

# Cover
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(LOGO),width=Inches(3.2)); p.paragraph_format.space_after=Pt(24)
para("קורסים נוספים לפרסום בהייטק זון",28,True,PURPLE,6,WD_ALIGN_PARAGRAPH.CENTER)
para("רפואה ויזמות  •  התנהלות כלכלית ושוק ההון",15,True,NAVY,22,WD_ALIGN_PARAGRAPH.CENTER)
para("שני קורסים מקוונים, מעשיים וחווייתיים לילדים ולבני נוער",11,False,GRAY,26,WD_ALIGN_PARAGRAPH.CENTER)
prices(); para("ימי שלישי | 17:00–18:30 | שישה מפגשים",12,True,NAVY,22,WD_ALIGN_PARAGRAPH.CENTER)
heading("פרטי הספק")
info([("שם","Z-SCHOOL"),("אופן הפעילות","אונליין — מכל הארץ"),("טלפון","053-3000045"),("דוא״ל","contact@z-school.co.il"),("אתר","www.z-school.co.il")])
para("הערת הפקה: מדיניות הביטולים תתווסף לאחר אישור הנוסח הסופי.",9.5,False,GRAY,0)

product(
    "החממה ליזמות רפואית-חברתית",
    "מכירים את גוף האדם, מתרגלים מיומנויות רפואיות ומפתחים רעיון שמשפר חיים",
    "medical-hitech-zone.jpg",
    "קורס מקוון ומעשי שמחבר בין רפואה, טכנולוגיה וחשיבה יזמית. המשתתפים חוקרים כיצד הגוף פועל, מתנסים באתגרים רפואיים בגובה העיניים, מזהים צורך אמיתי בקהילה ומפתחים עבורו פתרון בעל ערך.",
    [("קהל יעד","כיתות ד׳–י״ב, בתוכן המותאם לשלב הגיל"),("מתכונת","6 מפגשים אונליין, שעה וחצי בכל מפגש"),("יום ושעה","יום שלישי, 17:00–18:30"),("מועדים","13.10–17.11.2026"),("מיקום","Zoom דרך Z-Campus"),("פתיחת קבוצה","מינימום 4 משתתפים")],
    ["היכרות חווייתית עם גוף האדם, אנטומיה ועולם הרפואה.","תרגול עקרונות החייאה וקבלת משוב מהמדריך.","חשיפה לאבחון, טכנולוגיות רפואיות וכלי AI ברפואה.","זיהוי צורך רפואי או חברתי ופיתוח פתרון יזמי.","בניית קונספט, מיתוג והצגת המיזם בפיץ׳ מסכם.","למידה בהנחיה מקצועית ובתקשורת בגובה העיניים."],
    ["לא נדרש ידע מוקדם ברפואה או ביזמות.","נדרשים מחשב, מצלמה ואוזניות וחיבור אינטרנט יציב.","הפעילויות והדוגמאות מותאמות לגיל המשתתפים.","במקרה של היעדרות לא תישלח הקלטה."],
    "https://www.z-school.co.il/social-medical-entrepreneurship-course/"
)

product(
    "מאני מאסטר — התנהלות כלכלית ושוק ההון",
    "מבינים את חוקי הכסף ומקבלים כלים לעתיד כלכלי חכם",
    "finance-hitech-zone.jpg",
    "קורס חווייתי שמנגיש לילדים ולבני נוער את עולם הכסף בגובה העיניים. דרך דוגמאות, משימות ומשחקי החלטות לומדים כיצד לנהל תקציב, לקנות בצורה מושכלת, להבין בנקים וכרטיסי אשראי ולהכיר את עולם ההשקעות ושוק ההון.",
    [("קהל יעד","ילדים ובני נוער, בחלוקה לקבוצות גיל לפי ההרשמה"),("מתכונת","6 מפגשים אונליין, שעה וחצי בכל מפגש"),("יום ושעה","יום שלישי, 17:00–18:30"),("מועדים","13.10–17.11.2026"),("מיקום","Zoom דרך Z-Campus"),("פתיחת קבוצה","מינימום 4 משתתפים")],
    ["הבנת מושגים כלכליים והשפעתם על חיי היום-יום.","הכנסות, הוצאות, סדרי עדיפויות ובניית תקציב.","התנהלות מול בנקים וכרטיסי אשראי והימנעות מטעויות נפוצות.","צרכנות חכמה, השוואת מחירים וקבלת החלטות.","היכרות עם שוק ההון, מדדים, קריפטו ונדל״ן.","פיתוח חשיבה יזמית ותוכנית אישית לצעד כלכלי ראשון."],
    ["לא נדרש ידע מוקדם בכלכלה או בהשקעות.","נדרשים מחשב, מצלמה ואוזניות וחיבור אינטרנט יציב.","הקורס חינוכי ואינו מהווה ייעוץ השקעות.","במקרה של היעדרות לא תישלח הקלטה.","התוכן בהנחיית קורין חודר, בעלת תואר ראשון בכלכלה ומנכ״לית התאחדות היועצים הכלכליים."],
    "https://www.z-school.co.il/courses/economy-and-young-entrepreneurs/"
)

for sec in doc.sections:
    f=sec.footer; p=f.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in p.runs: r.text=""
    runfmt(p.add_run("Z-SCHOOL  |  תוכן מוצע לפרסום בהייטק זון  |  אוגוסט 2026"),8,False,GRAY)

doc.save(FINAL)
print("DOCX_CREATED")
