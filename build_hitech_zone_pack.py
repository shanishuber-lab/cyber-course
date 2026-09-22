from pathlib import Path
from PIL import Image, ImageEnhance
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"C:\Users\shani\Documents\course creator")
Image.MAX_IMAGE_PIXELS = None
OUT = ROOT / "output" / "hitech-zone"
IMG_DIR = OUT / "images"
OUT.mkdir(parents=True, exist_ok=True)
IMG_DIR.mkdir(parents=True, exist_ok=True)

SOURCES = {
    "english": Path(r"C:\Users\shani\.codex\generated_images\01a018d2-1ef6-7511-9026-5aa759e0d054\exec-27687a63-8cc4-4dfb-a53b-a9d5e717c53c.png"),
    "ai": Path(r"C:\Users\shani\.codex\generated_images\01a018d2-1ef6-7511-9026-5aa759e0d054\exec-eb9d0319-3211-477f-91f7-a4b4e1eaa381.png"),
    "cyber": Path(r"C:\Users\shani\.codex\generated_images\01a018d2-1ef6-7511-9026-5aa759e0d054\exec-6ca48c44-fd09-4689-973f-0ed726e11746.png"),
    "readers": Path(r"C:\Users\shani\.codex\generated_images\01a018d2-1ef6-7511-9026-5aa759e0d054\exec-7c175a5d-c5c3-4775-bfbb-64f3ff005430.png"),
}
LOGO = ROOT / "assets" / "logos" / "לוגו צבעוני על רקע שקוף.png"

def branded_image(src: Path, dst: Path):
    base = Image.open(src).convert("RGB")
    logo_rgb = Image.open(LOGO).convert("RGB")
    logo_rgb.thumbnail((2200, 1200), Image.Resampling.LANCZOS)
    # The supplied colored logo includes a very large white canvas. Crop to the
    # actual colored mark before placing it on the photo.
    mask = Image.new("L", logo_rgb.size)
    mask.putdata([255 if min(px) < 238 else 0 for px in logo_rgb.getdata()])
    bbox = mask.getbbox()
    logo_rgb = logo_rgb.crop(bbox)
    logo_mask = mask.crop(bbox)
    logo = logo_rgb.convert("RGBA")
    logo.putalpha(logo_mask)
    target_w = int(base.width * 0.145)
    ratio = target_w / logo.width
    logo = logo.resize((target_w, int(logo.height * ratio)), Image.Resampling.LANCZOS)
    # Create a quiet white badge so the original logo stays crisp on photography.
    pad_x, pad_y = 18, 12
    badge = Image.new("RGBA", (logo.width + 2 * pad_x, logo.height + 2 * pad_y), (255, 255, 255, 235))
    badge.alpha_composite(logo, (pad_x, pad_y))
    x, y = 28, 24
    canvas = base.convert("RGBA")
    canvas.alpha_composite(badge, (x, y))
    canvas.convert("RGB").save(dst, quality=95)

for key, src in SOURCES.items():
    branded_image(src, IMG_DIR / f"{key}-hitech-zone.jpg")

PURPLE = "5324F5"
NAVY = "07152B"
YELLOW = "FFD95A"
CYAN = "1AB8E8"
LIGHT = "F4F1FF"
GRAY = "595F6B"

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_margins(cell, top=120, start=150, bottom=120, end=150):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tcMar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")

def set_rtl(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.find(qn("w:bidi"))
    if bidi is None:
        bidi = OxmlElement("w:bidi")
        pPr.append(bidi)

def fmt_run(run, size=11, bold=False, color=NAVY):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().append(OxmlElement("w:rtl"))

def add_para(doc, text="", size=11, bold=False, color=NAVY, before=0, after=5, align=WD_ALIGN_PARAGRAPH.RIGHT):
    p = doc.add_paragraph()
    p.alignment = align
    if align == WD_ALIGN_PARAGRAPH.RIGHT:
        set_rtl(p)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    r = p.add_run(text)
    fmt_run(r, size=size, bold=bold, color=color)
    return p

def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    set_rtl(p)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text)
    fmt_run(r, size=10.5, color=NAVY)
    return p

def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    set_rtl(p)
    p.paragraph_format.space_before = Pt(9 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(text)
    fmt_run(r, size=18 if level == 1 else 13, bold=True, color=PURPLE if level == 1 else NAVY)
    return p

def add_info_table(doc, rows):
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for i, (label, value) in enumerate(rows):
        cells = table.rows[i].cells
        cells[0].width = Inches(4.7)
        cells[1].width = Inches(1.5)
        for c in cells:
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c)
        set_cell_shading(cells[1], LIGHT)
        p0 = cells[0].paragraphs[0]
        set_rtl(p0)
        fmt_run(p0.add_run(value), 10.5, color=NAVY)
        p1 = cells[1].paragraphs[0]
        set_rtl(p1)
        fmt_run(p1.add_run(label), 10.5, bold=True, color=PURPLE)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_price_strip(doc, regular, zone):
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    vals = [("מחיר רגיל", regular, LIGHT, PURPLE), ("מחיר הייטק זון", zone, YELLOW, NAVY)]
    for cell, (label, price, fill, color) in zip(table.rows[0].cells, vals):
        cell.width = Inches(3.1)
        set_cell_shading(cell, fill)
        set_cell_margins(cell, top=180, bottom=180)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fmt_run(p.add_run(label + "\n"), 10, bold=True, color=color)
        fmt_run(p.add_run(price), 18, bold=True, color=color)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)

def add_product(doc, title, subtitle, image_name, intro, facts, benefits, know, redemption, regular, zone, source=None):
    doc.add_page_break()
    p = add_para(doc, "Z-SCHOOL  |  הצעה להייטק זון", size=9, bold=True, color=CYAN, after=2)
    add_para(doc, title, size=24, bold=True, color=PURPLE, after=2)
    add_para(doc, subtitle, size=13, bold=True, color=GRAY, after=10)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.add_run().add_picture(str(IMG_DIR / image_name), width=Inches(6.45))
    p_img.paragraph_format.space_after = Pt(8)
    add_para(doc, intro, size=11, after=7)
    add_info_table(doc, facts)
    add_heading(doc, "מה מקבלים?", 2)
    for item in benefits:
        add_bullet(doc, item)
    add_heading(doc, "כדאי לדעת", 2)
    for item in know:
        add_bullet(doc, item)
    add_heading(doc, "אופן המימוש", 2)
    for i, item in enumerate(redemption, 1):
        add_para(doc, f"{i}. {item}", size=10.5, after=3)
    add_price_strip(doc, regular, zone)
    if source:
        add_para(doc, f"למידע נוסף: {source}", size=8.5, color=GRAY, after=0)

doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.65)
section.bottom_margin = Inches(0.65)
section.left_margin = Inches(0.8)
section.right_margin = Inches(0.8)
section.header_distance = Inches(0.25)
section.footer_distance = Inches(0.25)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
styles["Normal"]._element.rPr.rFonts.set(qn("w:cs"), "Arial")
styles["Normal"].font.size = Pt(11)

# Cover
p_logo = doc.add_paragraph()
p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_logo.add_run().add_picture(str(LOGO), width=Inches(3.2))
p_logo.paragraph_format.space_after = Pt(26)
add_para(doc, "חבילת תוכן לפרסום בהייטק זון", size=28, bold=True, color=PURPLE, after=6, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "ארבעה קורסים לילדים ולבני נוער", size=15, bold=True, color=NAVY, after=24, align=WD_ALIGN_PARAGRAPH.CENTER)
add_para(doc, "אנגלית מדוברת חווייתית  •  יצירת טריילר עם AI  •  CyberPath  •  קוראים חכמים", size=11, color=GRAY, after=30, align=WD_ALIGN_PARAGRAPH.CENTER)
add_price_strip(doc, "למידה מונחית: 800 ₪", "למידה מונחית: 600 ₪")
add_para(doc, "קורסים עצמאיים: מחיר רגיל 199 ₪ | מחיר הייטק זון 150 ₪", size=11, bold=True, color=NAVY, after=22, align=WD_ALIGN_PARAGRAPH.CENTER)
add_heading(doc, "פרטי הספק", 2)
add_info_table(doc, [
    ("שם", "Z-SCHOOL"),
    ("אופן הפעילות", "אונליין — מכל הארץ"),
    ("טלפון", "053-3000045"),
    ("דוא״ל", "contact@z-school.co.il"),
    ("אתר", "www.z-school.co.il"),
])
add_para(doc, "הערת הפקה: מדיניות הביטולים תתווסף לאחר אישור הנוסח הסופי.", size=9.5, color=GRAY, after=0)

add_product(
    doc,
    "אנגלית מדוברת — לומדים לדבר בביטחון",
    "קורס אונליין חווייתי ומשחקי לכיתות ג׳–ט׳",
    "english-hitech-zone.jpg",
    "שמונה מפגשים חיים שמוציאים את האנגלית מהמחברת ומכניסים אותה לשיחה. הילדים מתרגלים דיבור, הקשבה ותגובה באנגלית דרך משחקים, משימות ופעילויות קבוצתיות — בסביבה מעודדת שמחזקת את הביטחון ואת תחושת המסוגלות.",
    [
        ("קהל יעד", "שלוש קבוצות נפרדות: ג׳–ד׳, ה׳–ו׳, ז׳–ט׳"),
        ("מתכונת", "8 מפגשים אונליין, שעה בכל מפגש"),
        ("ימים ושעה", "ראשון ושלישי, 17:00–18:00"),
        ("מועד פתיחה", "יום שלישי, 13.10.2026"),
        ("מיקום", "Zoom דרך Z-Campus"),
        ("פתיחת קבוצה", "מינימום 4 משתתפים"),
    ],
    [
        "תרגול מעשי של שפה מדוברת בסיטואציות מחיי היום-יום.",
        "משחקי שפה, משימות קבוצתיות ופעילויות המעודדות השתתפות.",
        "חיזוק אוצר המילים, ההקשבה והיכולת להגיב באנגלית.",
        "קבוצה המותאמת לשכבת הגיל שנבחרה בעת ההרשמה.",
    ],
    [
        "לא נדרש ידע מוקדם מעבר ללימודי האנגלית המתאימים לשכבת הגיל.",
        "נדרשים מחשב, מצלמה ואוזניות וחיבור אינטרנט יציב.",
        "שלוש קבוצות הגיל מתקיימות במקביל.",
        "במקרה של היעדרות לא תישלח הקלטה.",
    ],
    [
        "לאחר הרכישה בהייטק זון מתקבל קוד הטבה.",
        "יוצרים קשר עם Z-SCHOOL ומציינים את הקוד ואת קבוצת הגיל המבוקשת.",
        "מקבלים קישור ייעודי, מזינים את הקוד ומשלימים את הרכישה במחיר ההטבה.",
        "פרטי הכניסה ל-Z-Campus ולמפגש נשלחים לאחר השלמת ההרשמה.",
    ],
    "800 ₪", "600 ₪"
)

add_product(
    doc,
    "יוצרים טריילר עם בינה מלאכותית",
    "קורס מונחה ב-LUMO לכיתות ד׳–ו׳ — מרעיון ועד טריילר אישי",
    "ai-hitech-zone.jpg",
    "בקורס יצירתי בן שישה מפגשים כל ילד וילדה מפתחים רעיון, בונים סיפור ויוצרים טריילר אישי בעזרת בינה מלאכותית. העבודה מתבצעת בתוך LUMO — סביבת יצירה דיגיטלית מסודרת שמרכזת את כלי ה-AI ואת גלריית התוצרים במקום אחד.",
    [
        ("קהל יעד", "כיתות ד׳–ו׳"),
        ("מתכונת", "6 מפגשים אונליין, שעה וחצי בכל מפגש"),
        ("יום ושעה", "יום שני, 17:00–18:30"),
        ("מועדים", "12.10–16.11.2026"),
        ("מיקום", "Zoom דרך Z-Campus"),
        ("פתיחת קבוצה", "מינימום 4 משתתפים"),
    ],
    [
        "פיתוח רעיון וסיפור וכתיבת תסריט בעזרת AI.",
        "יצירת דמויות ותמונות המתאימות לעולם הסיפור.",
        "הנפשת תמונות ויצירת קטעי וידאו.",
        "הוספת מוזיקה, קול ואפקטים.",
        "עריכה והצגת טריילר אישי כפרויקט גמר.",
        "חשבון LUMO וקרדיטים ליצירה כלולים במחיר.",
    ],
    [
        "לא נדרש ידע מוקדם בבינה מלאכותית או בעריכת וידאו.",
        "נדרשים מחשב, מצלמה ואוזניות וחיבור אינטרנט יציב.",
        "חשבון LUMO פעיל לאורך הקורס בלבד.",
        "במקרה של היעדרות לא תישלח הקלטה.",
    ],
    [
        "לאחר הרכישה בהייטק זון מתקבל קוד הטבה.",
        "יוצרים קשר עם Z-SCHOOL ומציינים את הקוד.",
        "מקבלים קישור ייעודי, מזינים את הקוד ומשלימים רכישה במחיר ההטבה.",
        "חשבון LUMO ופרטי הכניסה לקורס נשלחים לפני תחילת הפעילות.",
    ],
    "800 ₪ כולל קרדיטים", "600 ₪ כולל קרדיטים",
    "https://guide.lumoai.co.il/"
)

add_product(
    doc,
    "CyberPath — קורס סייבר לנוער",
    "לומדים לחשוב כמו אנשי סייבר — בסביבה בטוחה ובקצב אישי",
    "cyber-hitech-zone.jpg",
    "קורס דיגיטלי מעשי שמלמד איך רשתות מחשבים עובדות, כיצד מזהים את המכשירים והשירותים ברשת ואיך מבינים את הצד ההתקפי כדי לבנות הגנה טובה יותר. כל יחידה משלבת וידאו קצר, סימולטור ותרגול עם משוב מיידי.",
    [
        ("קהל יעד", "גילאי 13–18"),
        ("מתכונת", "למידה עצמית — 8 יחידות"),
        ("משך יחידה", "כ-20–35 דקות של למידה פעילה"),
        ("ידע קודם", "לא נדרש"),
        ("גישה", "ללא הגבלת זמן, כולל עדכונים עתידיים"),
        ("מכשירים", "מחשב, טאבלט או טלפון דרך הדפדפן"),
    ],
    [
        "היכרות עם רשתות, כתובות IP ונתבים.",
        "סריקת רשת, מיפוי מכשירים והבנת פורטים ושירותים.",
        "עבודה עם כלים ומושגים מעולם הסייבר בסביבה בטוחה.",
        "סימולטורי טרמינל, תרגולים ומשוב מיידי.",
        "בוחן סיום ואישור סיום.",
    ],
    [
        "הקורס בנוי מאפס ומתאים גם למי שלא התנסה בתכנות.",
        "הלמידה עצמאית וגמישה, ללא מפגשים חיים.",
        "אין צורך בהתקנת תוכנה.",
    ],
    [
        "לאחר הרכישה בהייטק זון מתקבל קוד הטבה.",
        "יוצרים קשר עם Z-SCHOOL ומציינים את הקוד.",
        "מקבלים קישור ייעודי, מזינים את הקוד ומשלימים את הרכישה במחיר ההטבה.",
        "לאחר השלמת הרכישה מתקבל קישור ישיר לקמפוס ולתוכן הקורס.",
    ],
    "199 ₪", "150 ₪",
    "https://cyber.z-school.co.il/"
)

add_product(
    doc,
    "קוראים חכמים — הבנת הנקרא שהופכת להרפתקה",
    "עולם משחק דיגיטלי לתלמידי כיתות ג׳–ד׳",
    "readers-hitech-zone.jpg",
    "קורס עצמאי שהופך את הבנת הנקרא למסע משחקי: הילדים מתקדמים במפה, מגדלים דרקון קריאה, צדים מילים, פותרים חידות וכותבים בעצמם. כל שלב בונה ביטחון בהדרגה — ממילים ומשפטים ועד קטע מלא והסקת מסקנות.",
    [
        ("קהל יעד", "תלמידי כיתות ג׳–ד׳"),
        ("מתכונת", "למידה עצמית — 8 יחידות הרפתקה"),
        ("משך יחידה", "כ-20–30 דקות"),
        ("משחקים", "12 סוגי משחקונים עם משוב מיידי"),
        ("מכשירים", "מחשב או טאבלט דרך הדפדפן"),
        ("תשלום", "חד-פעמי, ללא מנוי"),
    ],
    [
        "תרגול הבנת הנקרא בדרך הדרגתית ומשחקית.",
        "קריאה עם ניקוד מלא לפי הצורך, או ללא ניקוד.",
        "שאלות הבנה מהמידע המפורש ועד הסקה.",
        "תרגול כתיבה — מכותרת ועד משפטים שלמים.",
        "דרקון אישי, מטבעות, שדרוגים וספר מסע.",
    ],
    [
        "התוכן נבנה עם אנשי חינוך ובהתאם לתוכנית הלימודים.",
        "טעויות מקבלות משוב מסביר ומעודד.",
        "הקורס מתאים במיוחד לילדים שדפי עבודה רגילים אינם מניעים אותם.",
        "אין צורך בהתקנה.",
    ],
    [
        "לאחר הרכישה בהייטק זון מתקבל קוד הטבה.",
        "יוצרים קשר עם Z-SCHOOL ומציינים את הקוד.",
        "מקבלים קישור ייעודי, מזינים את הקוד ומשלימים את הרכישה במחיר ההטבה.",
        "לאחר השלמת הרכישה מתקבל קישור ישיר לקמפוס ולתוכן הקורס.",
    ],
    "199 ₪", "150 ₪",
    "https://www.z-school.co.il/smart-readers/"
)

# Footer
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Z-SCHOOL  |  תוכן מוצע לפרסום בהייטק זון  |  אוגוסט 2026")
    fmt_run(r, size=8, color=GRAY)

output_path = OUT / "חבילת תוכן להייטק זון - Z-SCHOOL.docx"
doc.save(output_path)
print("DOCX_CREATED")
